import os
import random
import qrcode
import string
import matplotlib.pyplot as plt
from pyzbar.pyzbar import decode, ZBarSymbol
from PIL import ImageDraw
from multiprocessing import Pool, cpu_count
from docx import Document

cover_mode = "Logo" # choose between "RandomBits", "Logo", "Side", "Border", "RandomBytes", "BytesInOrder" and "Maximum"
max_attempts = 100 # only for cover modes "RandomBits" and "RandomBytes"

cover_color = "blue"
box_size = 4
border = 4

error_correction_levels = {
    "L": [qrcode.constants.ERROR_CORRECT_L, 0.07],
    "M": [qrcode.constants.ERROR_CORRECT_M, 0.15],
    "Q": [qrcode.constants.ERROR_CORRECT_Q, 0.25],
    "H": [qrcode.constants.ERROR_CORRECT_H, 0.30]
}

alignment_pattern_positions = {
    1: [],
    2: [6, 18],
    3: [6, 22],
    4: [6, 26],
    5: [6, 30],
    6: [6, 34],
    7: [6, 22, 38],
    8: [6, 24, 42],
    9: [6, 26, 46],
    10: [6, 28, 50],
    11: [6, 30, 54],
    12: [6, 32, 58],
    13: [6, 34, 62],
    14: [6, 26, 46, 66],
    15: [6, 26, 48, 70],
    16: [6, 26, 50, 74],
    17: [6, 30, 54, 78],
    18: [6, 30, 56, 82],
    19: [6, 30, 58, 86],
    20: [6, 34, 62, 90],
    21: [6, 28, 50, 72, 94],
    22: [6, 26, 50, 74, 98],
    23: [6, 30, 54, 78, 102],
    24: [6, 28, 54, 80, 106],
    25: [6, 32, 58, 84, 110],
    26: [6, 30, 58, 86, 114],
    27: [6, 34, 62, 90, 118],
    28: [6, 26, 50, 74, 98, 122],
    29: [6, 30, 54, 78, 102, 126],
    30: [6, 26, 52, 78, 104, 130],
    31: [6, 30, 56, 82, 108, 134],
    32: [6, 34, 60, 86, 112, 138],
    33: [6, 30, 58, 86, 114, 142],
    34: [6, 34, 62, 90, 118, 146],
    35: [6, 30, 54, 78, 102, 126, 150],
    36: [6, 24, 50, 76, 102, 128, 154],
    37: [6, 28, 54, 80, 106, 132, 158],
    38: [6, 32, 58, 84, 110, 136, 162],
    39: [6, 26, 54, 82, 110, 138, 166],
    40: [6, 30, 58, 86, 114, 142, 170]
}

max_bytes = {
    1:  {"L":17,"M":14,"Q":11,"H":7},
    2:  {"L":32,"M":26,"Q":20,"H":14},
    3:  {"L":53,"M":42,"Q":32,"H":24},
    4:  {"L":78,"M":62,"Q":46,"H":34},
    5:  {"L":106,"M":84,"Q":60,"H":44},
    6:  {"L":134,"M":106,"Q":74,"H":58},
    7:  {"L":154,"M":122,"Q":86,"H":64},
    8:  {"L":192,"M":152,"Q":108,"H":84},
    9:  {"L":230,"M":180,"Q":130,"H":98},
    10: {"L":271,"M":213,"Q":151,"H":119},
    11: {"L":321,"M":251,"Q":177,"H":137},
    12: {"L":367,"M":287,"Q":203,"H":155},
    13: {"L":425,"M":331,"Q":241,"H":177},
    14: {"L":458,"M":362,"Q":258,"H":194},
    15: {"L":520,"M":412,"Q":292,"H":220},
    16: {"L":586,"M":450,"Q":322,"H":250},
    17: {"L":644,"M":504,"Q":364,"H":280},
    18: {"L":718,"M":560,"Q":394,"H":310},
    19: {"L":792,"M":624,"Q":442,"H":338},
    20: {"L":858,"M":666,"Q":482,"H":382},
    21: {"L":929,"M":711,"Q":509,"H":403},
    22: {"L":1003,"M":779,"Q":565,"H":439},
    23: {"L":1091,"M":857,"Q":611,"H":461},
    24: {"L":1171,"M":911,"Q":661,"H":511},
    25: {"L":1273,"M":997,"Q":715,"H":535},
    26: {"L":1367,"M":1059,"Q":751,"H":593},
    27: {"L":1465,"M":1125,"Q":805,"H":625},
    28: {"L":1528,"M":1190,"Q":868,"H":658},
    29: {"L":1628,"M":1264,"Q":908,"H":698},
    30: {"L":1732,"M":1370,"Q":982,"H":742},
    31: {"L":1840,"M":1452,"Q":1030,"H":790},
    32: {"L":1952,"M":1538,"Q":1112,"H":842},
    33: {"L":2068,"M":1628,"Q":1168,"H":898},
    34: {"L":2188,"M":1722,"Q":1228,"H":958},
    35: {"L":2303,"M":1809,"Q":1283,"H":983},
    36: {"L":2431,"M":1911,"Q":1351,"H":1051},
    37: {"L":2563,"M":1989,"Q":1423,"H":1093},
    38: {"L":2699,"M":2099,"Q":1499,"H":1139},
    39: {"L":2809,"M":2213,"Q":1579,"H":1219},
    40: {"L":2953,"M":2331,"Q":1663,"H":1273},
}

def cover_qr(args):
    level, version = args

    data = "jfsg.nl"
    data_bytes = max_bytes[version][level]
    if data_bytes > len(data):
        data += "/" + "".join(
            random.choice(string.ascii_letters + string.digits)
            for _ in range(data_bytes - len(data) - 1)
        )

    qr = qrcode.QRCode(
        error_correction=error_correction_levels[level][0],
        box_size=box_size,
        border=border
    )
    qr.add_data(data)
    base_img = qr.make_image().convert("RGB")
    filename = os.path.join(cover_mode, "qrcodes", f"qrcode{level}{version}.png")

    grid_size = base_img.size[0] // box_size - 2 * border

    def is_function_pattern(x, y):
        if x < 8 and y < 8:
            return True
        if x >= grid_size - 8 and y < 8:
            return True
        if x < 8 and y >= grid_size - 8:
            return True
        if x == 6 or y == 6:
            return True
        if x == 8 and y == grid_size - 8:
            return True

        for cx in alignment_pattern_positions[version]:
            for cy in alignment_pattern_positions[version]:
                if (cx < 9 and cy < 9) or (cx >= grid_size - 8 and cy < 8) or (cx < 8 and cy >= grid_size - 8):
                    continue
                if abs(x - cx) <= 2 and abs(y - cy) <= 2:
                    return True
        return False

    def is_format_information(x, y):
        if x == 8 and (y < 6 or y == 7):
            return True
        if (x < 6 or 7 <= x <= 8) and y == 8:
            return True
        if x >= grid_size - 8 and y == 8:
            return True
        if x == 8 and y > grid_size - 8:
            return True
        return False

    def is_version_info(x, y):
        if version < 7:
            return False
        if grid_size - 11 <= x <= grid_size - 9 and y < 6:
            return True
        if x < 6 and grid_size - 11 <= y <= grid_size - 9:
            return True
        return False

    qr_bytes = [[]]
    for p in range(grid_size * (grid_size - 1) + 8):
        x = (grid_size - 1) - p // (2 * grid_size) * 2 - (p % 2)
        y = p % (2 * grid_size) // 2
        if x <= 6:
            x -= 1
        if p // (grid_size * 2) % 2 == 0:
            y = grid_size - 1 - y
        if is_function_pattern(x, y) or is_format_information(x, y) or is_version_info(x, y):
            continue
        x1 = (x + border) * box_size
        y1 = (y + border) * box_size
        x2 = x1 + box_size - 1
        y2 = y1 + box_size - 1

        byte = qr_bytes[len(qr_bytes) - 1]
        if len(byte) == 8:
            qr_bytes.append([(x1, y1, x2, y2)])
        else:
            byte.append((x1, y1, x2, y2))

    if cover_mode == "RandomBits":
        attempts = 0
        cover_percentage = 0.005
        readable_img = base_img.copy()

        cover_bits = []
        for byte in qr_bytes:
            for bit in byte:
                cover_bits.append(bit)

        while True:
            img = base_img.copy()
            draw = ImageDraw.Draw(img)
            covered_bytes = []
            for rect in random.sample(cover_bits, int(len(qr_bytes) * 8 * cover_percentage)):
                px = img.load()
                x1, y1, x2, y2 = rect
                cx = (x1 + x2) // 2
                cy = (y1 + y2) // 2

                r, g, b = px[cx, cy]
                is_black = (r + g + b) == 0
                new_color = (255, 255, 0) if is_black else (0, 0, 255)
                draw.rectangle(rect, fill=new_color)
                for byte in qr_bytes:
                    for bit in byte:
                        if rect == bit and byte not in covered_bytes:
                            covered_bytes.append(byte)
                if len(covered_bytes) / len(qr_bytes) >= cover_percentage:
                    break
            attempts += 1

            if decode(img, symbols=[ZBarSymbol.QRCODE]):
                attempts = 0
                cover_percentage += 0.01
                readable_img = img.copy()
                continue

            if attempts >= max_attempts:
                readable_img.save(filename)
                return level, version, cover_percentage - 0.01

    elif cover_mode == "Logo" or cover_mode == "Side" or cover_mode == "Border":
        size = 1
        pixels_to_cover = 0
        if cover_mode == "Logo":
            pixels_to_cover = size * size
        elif cover_mode == "Side":
            pixels_to_cover = size * (size + 1) // 2
        elif cover_mode == "Border":
            pixels_to_cover = 4 * size * (grid_size - size)
        cover_percentage = 0
        readable_img = base_img.copy()
        max_cover = 0

        while True:
            cover_pixels = []
            for p in range(pixels_to_cover):
                x = 0
                y = 0
                if cover_mode == "Logo":
                    x = p % size + (grid_size - size + 1) // 2
                    y = p // size + (grid_size - size + 1) // 2
                elif cover_mode == "Side":
                    y = int(((8 * p + 1) ** 0.5 - 1) // 2)
                    x = size - 1 - y + p - y * (y + 1) // 2 + grid_size - size
                    y += grid_size - size
                elif cover_mode == "Border":
                    if p < grid_size * size:
                        x = p % grid_size
                        y = p // grid_size
                    elif grid_size * size <= p < pixels_to_cover - grid_size * size:
                        x = p - grid_size * size - (2 * size) * ((p - grid_size * size) // (2 * size))
                        if (p - grid_size * size) % (2 * size) >= size:
                            x += grid_size - 2 * size
                        y = size + (p - grid_size * size) // (2 * size)
                    else:
                        x = (grid_size - (pixels_to_cover - p)) % grid_size
                        y = (grid_size - (pixels_to_cover - p)) // grid_size + grid_size - 1
                x1 = (x + border) * box_size
                y1 = (y + border) * box_size
                x2 = x1 + box_size - 1
                y2 = y1 + box_size - 1
                cover_pixels.append((x1, y1, x2, y2))

            img = base_img.copy()
            draw = ImageDraw.Draw(img)
            covered_bytes = []
            for rect in cover_pixels:
                px = img.load()
                x1, y1, x2, y2 = rect
                cx = (x1 + x2) // 2
                cy = (y1 + y2) // 2

                r, g, b = px[cx, cy]
                is_black = (r + g + b) == 0
                if (rect[0] == 80 or rect[1] == 80) and version == 1 and size == 5:
                    print(is_black, version, rect)
                draw.rectangle(rect, fill=cover_color)
                if not is_black:
                    for byte in qr_bytes:
                        for bit in byte:
                            if rect == bit and byte not in covered_bytes:
                                covered_bytes.append(byte)

            cover_percentage = len(covered_bytes) / len(qr_bytes)

            if decode(img, symbols=[ZBarSymbol.QRCODE]):
                readable_img = img.copy()
                max_cover = cover_percentage
            elif cover_percentage >= error_correction_levels[level][1] + 0.05:
                readable_img.save(filename)
                return level, version, max_cover

            if cover_mode == "Logo":
                size += 2
                pixels_to_cover = size * size
            elif cover_mode == "Side":
                size += 1
                pixels_to_cover = size * (size + 1) // 2
            elif cover_mode == "Border":
                size += 1
                pixels_to_cover = 4 * size * (grid_size - size)

    elif cover_mode == "RandomBytes":
        attempts = 0
        bytes_to_cover = 1
        readable_img = base_img.copy()
        max_cover = 0

        while True:
            img = base_img.copy()
            draw = ImageDraw.Draw(img)

            random_bytes = random.sample(qr_bytes, bytes_to_cover)
            random_bits = []
            for byte in random_bytes:
                for bit in byte:
                    random_bits.append(bit)

            for rect in random_bits:
                draw.rectangle(rect, fill=cover_color)

            attempts += 1

            if decode(img, symbols=[ZBarSymbol.QRCODE]):
                attempts = 0
                bytes_to_cover += 1
                readable_img = img.copy()
                max_cover = len(random_bytes) / len(qr_bytes)
                continue

            if attempts >= max_attempts:
                readable_img.save(filename)
                return level, version, max_cover

    elif cover_mode == "BytesInOrder":
        bytes_to_cover = 1
        readable_img = base_img.copy()

        while True:
            img = base_img.copy()
            draw = ImageDraw.Draw(img)

            random_bits = []
            for byte in qr_bytes[:bytes_to_cover]:
                for bit in byte:
                    random_bits.append(bit)

            for rect in random_bits:
                draw.rectangle(rect, fill=cover_color)

            if decode(img, symbols=[ZBarSymbol.QRCODE]):
                bytes_to_cover += 1
                readable_img = img.copy()
            else:
                readable_img.save(filename)
                return level, version, (bytes_to_cover - 1) / len(qr_bytes)

    elif cover_mode == "Maximum":
        total_bytes = len(qr_bytes)
        data_bytes = max_bytes[version][level]
        ecc_bytes = len(qr_bytes) - data_bytes - (2 if version <= 9 else 3)
        error_correction = (ecc_bytes // 2) / total_bytes
        return level, version, error_correction

    return level, version, 0

if __name__ == "__main__":
    if os.path.exists(cover_mode):
        for x in os.listdir(cover_mode):
            if os.path.isdir(os.path.join(cover_mode, x)):
                for y in os.listdir(os.path.join(cover_mode, x)):
                    os.remove(os.path.join(cover_mode, x, y))
            else:
                os.remove(os.path.join(cover_mode, x))

    os.makedirs(cover_mode, exist_ok=True)
    os.makedirs(os.path.join(cover_mode, "qrcodes"), exist_ok=True)

    tasks = [(level, version)
             for level in error_correction_levels
             for version in range(1, 41)]

    doc = Document()
    table = doc.add_table(rows=41, cols=5)
    table.style = "Table Grid"

    headers = ["Versie", "L", "M", "Q", "H"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    data = [[i, 0, 0, 0, 0] for i in range(1, 41)]

    with Pool(max(1, cpu_count() // 2)) as pool:
        for level, version, cover_percentage in pool.imap(cover_qr, tasks, chunksize=1):
            data[version - 1][headers.index(level)] = cover_percentage
            print(f"QR {level} v{version} was covered for {round(cover_percentage * 100)}%")

    versions = [row[0] for row in data]
    L_vals = [row[1] * 100 for row in data]
    M_vals = [row[2] * 100 for row in data]
    Q_vals = [row[3] * 100 for row in data]
    H_vals = [row[4] * 100 for row in data]

    plt.figure(figsize=(10, 6))

    plt.plot(versions, L_vals, label="L")
    plt.plot(versions, M_vals, label="M")
    plt.plot(versions, Q_vals, label="Q")
    plt.plot(versions, H_vals, label="H")

    plt.xlabel("Versie QR-code (1-40)")
    plt.ylabel("Maximaal percentage afgedekte bits (%)")
    plt.title("Maximale afdekking per versie QR-code")
    plt.ylim(0, 100)

    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    plt.savefig(os.path.join(cover_mode, "graph.png"), dpi=300)
    plt.show()

    for version, row in enumerate(data, start=1):
        for col, value in enumerate(row):
            if col == 0:
                table.cell(version, col).text = str(value)
            else:
                table.cell(version, col).text = f"{round(value * 100)}%"

    doc.save(os.path.join(cover_mode, "table.docx"))
    print("Done")